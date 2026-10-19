from django.contrib import admin
from django.utils.html import format_html
from django.utils import timezone
from .models import Client, Ticket, Comment, ImportLog
from .admin_views import excel_import_admin_view

class CommentInline(admin.TabularInline):
    model = Comment
    extra = 1
    fields = ('author', 'text', 'created_at')
    readonly_fields = ('created_at',)
    verbose_name = "Внутренний комментарий"
    verbose_name_plural = "Внутренние комментарии по заявке"


@admin.register(Ticket)
class TicketAdmin(admin.ModelAdmin):
    list_display = (
        'external_id', 'subject', 'client', 
        'colored_priority', 'colored_status', 
        'assignee', 'created_at'
    )
    list_filter = ('priority', 'status', 'assignee', 'created_at')
    search_fields = ('subject', 'client__name', 'external_id')
    inlines = [CommentInline]
    readonly_fields = ('created_at', 'closed_at')
    actions = ['assign_to_me', 'close_tickets', 'export_selected']

    # подключаем кнопку на панель 
    change_list_template = "admin/HelpDesk/ticket/change_list.html"

    # путь от одного файла в другой
    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path('excel-import/', self.admin_site.admin_view(lambda r: excel_import_admin_view(self, r)), name='ticket_excel_import'),
        ]
        return custom_urls + urls

    # цветные бейджы кнопок
    @admin.display(description="Приоритет")
    def colored_priority(self, obj):
        colors = {
            Ticket.Priority.LOW: '#93c5fd', Ticket.Priority.NORMAL: '#10b981',
            Ticket.Priority.HIGH: '#f59e0b', Ticket.Priority.CRITICAL: '#ef4444',
        }
        color = colors.get(obj.priority, '#6b7280')
        return format_html('<span style="background-color: {}; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;">{}</span>', color, obj.get_priority_display())

    @admin.display(description="Статус")
    def colored_status(self, obj):
        colors = {
            Ticket.Status.NEW: '#3b82f6', Ticket.Status.IN_PROGRESS: '#8b5cf6',
            Ticket.Status.WAITING: '#f59e0b', Ticket.Status.CLOSED: '#6b7280',
        }
        color = colors.get(obj.status, '#374151')
        return format_html('<span style="background-color: {}; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;">{}</span>', color, obj.get_status_display())

        # сохранении модели через админку 
    def save_model(self, request, obj, form, change):
        # если заявка создается вручную ставим текущее время
        if not obj.created_at:
            obj.created_at = timezone.now()

        # если статус меняется на закрыт, ставим время закрытия 
        if obj.status == Ticket.Status.CLOSED and not obj.closed_at:
            obj.closed_at = timezone.now()
        # если статус изменили, сбрасываем дату закрытия
        elif obj.status != Ticket.Status.CLOSED:
            obj.closed_at = None

        super().save_model(request, obj, form, change)


    # --- ACTIONS ПО ТЗ ---
    @admin.action(description="Назначить выбранные заявки на меня")
    def assign_to_me(self, request, queryset):
        updated_count = queryset.update(assignee=request.user)
        self.message_user(request, f"Успешно: {updated_count} заявок назначено на вас.")

    @admin.action(description="Массово закрыть выбранные заявки")
    def close_tickets(self, request, queryset):
        open_tickets = queryset.exclude(status=Ticket.Status.CLOSED)
        updated_count = open_tickets.update(status=Ticket.Status.CLOSED, closed_at=timezone.now())
        self.message_user(request, f"Успешно: {updated_count} заявок массово закрыто.")

    @admin.action(description="Выгрузить выбранные заявки в Excel (.xlsx)")
    def export_selected(self, request, queryset):
        import openpyxl
        from django.http import HttpResponse
        from openpyxl.styles import PatternFill, Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Выгрузка заявок"
        ws.append(["ID заявки", "Тема", "Клиент", "Приоритет", "Статус", "Исполнитель", "Дата создания"])
        
        header_fill = PatternFill(start_color="1e1e2d", end_color="1e1e2d", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        critical_fill = PatternFill(start_color="ffeaec", end_color="ffeaec", fill_type="solid")
        critical_font = Font(name="Arial", size=11, color="990000", bold=True)
        
        for ticket in queryset:
            ws.append([
                ticket.external_id, ticket.subject, ticket.client.name,
                ticket.get_priority_display(), ticket.get_status_display(),
                ticket.assignee.username if ticket.assignee else "Не назначен",
                ticket.created_at.strftime("%Y-%m-%d %H:%M") if ticket.created_at else ""
            ])
            if ticket.priority == Ticket.Priority.CRITICAL:
                for cell in ws[ws.max_row]:
                    cell.fill = critical_fill
                    cell.font = critical_font
                    
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col.column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f"attachment; filename=tickets_export_{timezone.now().strftime('%Y%m%d')}.xlsx"
        wb.save(response)
        return response


@admin.register(Client)
class ClientAdmin(admin.ModelAdmin):
    list_display = ('name', 'email', 'phone', 'company', 'created_at')
    search_fields = ('name', 'email', 'company')


@admin.register(ImportLog)
class ImportLogAdmin(admin.ModelAdmin):
    list_display = ('file_name', 'user', 'uploaded_at', 'total_rows', 'created_count', 'updated_count', 'skipped_count')
    list_filter = ('uploaded_at', 'user')
    readonly_fields = ('user', 'file_name', 'uploaded_at', 'total_rows', 'created_count', 'updated_count', 'skipped_count', 'errors')
